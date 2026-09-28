from pathlib import Path
import hashlib
from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from xml.sax.saxutils import escape

ROOT=Path(__file__).resolve().parents[1]
ARTICLE=ROOT/"article"/"ARTICLE_FOR_COPYING.txt"
SOURCES=ROOT/"evidence"/"SOURCES_AND_CALCULATIONS.txt"
HEADINGS={
"Where the 228 billion came from","Whose water are we talking about?","The backup is part of the proposal",
"Build something that fits the source","Those 22 minutes belong to somebody","What did it replace?",
"I work on these systems. Check mine too.","Who agreed to the bill?","The restriction needs a reason too","Where I land"
}

def blocks(text):
    out=[]; cur=[]
    for line in text.replace("\r","").split("\n"):
        s=line.strip()
        if not s:
            if cur: out.append(("p"," ".join(cur))); cur=[]
        elif s in HEADINGS:
            if cur: out.append(("p"," ".join(cur))); cur=[]
            out.append(("h",s))
        else: cur.append(s)
    if cur: out.append(("p"," ".join(cur)))
    return out

def build_docx():
    text=ARTICLE.read_text(encoding="utf-8")
    parts=blocks(text)
    doc=Document()
    sec=doc.sections[0]; sec.top_margin=Inches(.75); sec.bottom_margin=Inches(.75)
    sec.left_margin=Inches(.85); sec.right_margin=Inches(.85)
    normal=doc.styles["Normal"]; normal.font.name="Georgia"; normal.font.size=Pt(11)
    for i,(k,s) in enumerate(parts):
        if i==0:
            p=doc.add_paragraph(); r=p.add_run(s); r.bold=True; r.font.size=Pt(24)
        elif i==1:
            p=doc.add_paragraph(s); p.style=doc.styles["Subtitle"]
        elif k=="h": doc.add_heading(s,level=1)
        else: doc.add_paragraph(s)
    doc.save(ROOT/"article"/"About_That_228_Billion_Gallons.docx")

def page_num(canvas,doc):
    canvas.saveState(); canvas.setFont("Helvetica",8); canvas.drawCentredString(LETTER[0]/2,.42*inch,str(doc.page)); canvas.restoreState()

def make_pdf(src,out,title):
    text=src.read_text(encoding="utf-8")
    styles=getSampleStyleSheet()
    body=ParagraphStyle("body",parent=styles["BodyText"],fontName="Times-Roman",fontSize=10.6,leading=15.2,spaceAfter=9)
    h=ParagraphStyle("h",parent=styles["Heading2"],fontName="Helvetica-Bold",fontSize=15,leading=18,spaceBefore=14,spaceAfter=8)
    title_s=ParagraphStyle("title",parent=styles["Title"],fontName="Helvetica-Bold",fontSize=22,leading=25,alignment=TA_CENTER,spaceAfter=12)
    by=ParagraphStyle("by",parent=styles["BodyText"],fontName="Helvetica",fontSize=9,leading=12,alignment=TA_CENTER,spaceAfter=18)
    story=[]; parts=blocks(text)
    for i,(k,s) in enumerate(parts):
        if i==0: story.append(Paragraph(escape(s),title_s))
        elif i==1: story.append(Paragraph(escape(s),by))
        elif k=="h": story.append(Paragraph(escape(s),h))
        else: story.append(Paragraph(escape(s),body))
    SimpleDocTemplate(str(out),pagesize=LETTER,rightMargin=.72*inch,leftMargin=.72*inch,topMargin=.7*inch,bottomMargin=.65*inch,title=title,author="Robert Thomas King").build(story,onFirstPage=page_num,onLaterPages=page_num)

def sources_pdf():
    text=SOURCES.read_text(encoding="utf-8")
    styles=getSampleStyleSheet()
    body=ParagraphStyle("sbody",parent=styles["BodyText"],fontName="Helvetica",fontSize=8.6,leading=12,spaceAfter=7)
    head=ParagraphStyle("shead",parent=styles["Heading2"],fontName="Helvetica-Bold",fontSize=13,leading=16,spaceBefore=12,spaceAfter=7)
    title=ParagraphStyle("stitle",parent=styles["Title"],fontName="Helvetica-Bold",fontSize=20,leading=23,spaceAfter=14)
    story=[]; paras=[p.strip() for p in text.replace("\r","").split("\n\n") if p.strip()]
    for i,p in enumerate(paras):
        style=title if i==0 else head if (len(p)<100 and not p.startswith("http")) else body
        story.append(Paragraph(escape(p).replace("\n","<br/>"),style))
    SimpleDocTemplate(str(ROOT/"evidence"/"Sources_and_Calculations.pdf"),pagesize=LETTER,rightMargin=.65*inch,leftMargin=.65*inch,topMargin=.65*inch,bottomMargin=.65*inch,title="Sources and calculations").build(story,onFirstPage=page_num,onLaterPages=page_num)

def manifest():
    paths=[
      ROOT/"article"/"ARTICLE_FOR_COPYING.txt",ROOT/"article"/"About_That_228_Billion_Gallons.docx",
      ROOT/"article"/"About_That_228_Billion_Gallons.pdf",ROOT/"evidence"/"SOURCES_AND_CALCULATIONS.txt",
      ROOT/"evidence"/"SOURCE_REGISTER.json",ROOT/"evidence"/"Sources_and_Calculations.pdf",
      ROOT/"analysis"/"reproduce_calculations.py",ROOT/"index.html"
    ]
    lines=[]
    for p in paths:
        if p.exists(): lines.append(f"{hashlib.sha256(p.read_bytes()).hexdigest()}  {p.relative_to(ROOT)}")
    (ROOT/"RELEASE_SHA256.txt").write_text("\n".join(lines)+"\n",encoding="utf-8")

if __name__=="__main__":
    build_docx()
    make_pdf(ARTICLE,ROOT/"article"/"About_That_228_Billion_Gallons.pdf","About That 228 Billion Gallons")
    sources_pdf()
    manifest()
    print("Built publication artifacts and RELEASE_SHA256.txt")
